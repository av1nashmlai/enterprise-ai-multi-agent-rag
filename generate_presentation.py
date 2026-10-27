"""
Generates a professional 2-slide PowerPoint presentation (architecture_presentation.pptx)
strictly satisfying the Edureka! / Illinois Tech Evaluator Criteria:
- Slide 1: System Architecture & Autonomous Multi-Agent Workflow (Planner -> Retriever -> Reasoner -> Validator)
- Slide 2: Ingestion Pipeline, Reliability Controls, and Key Limitations & Challenges
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN


def build_presentation(output_path="architecture_presentation.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # --- SLIDE 1: Architecture & Multi-Agent Workflow ---
    slide1 = prs.slides.add_slide(blank_layout)

    # Title Banner
    tx_box = slide1.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.2))
    tf = tx_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Enterprise AI Knowledge & Decision Support System"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(15, 23, 42)
    p2 = tf.add_paragraph()
    p2.text = "Generative AI & ML Capstone Project | Illinois Tech & Edureka! — Autonomous Multi-Agent RAG Architecture"
    p2.font.size = Pt(14)
    p2.font.color.rgb = RGBColor(71, 85, 105)

    # 4 Agent Columns
    agents = [
        ("1. Planner Agent", "Query Decomposition", [
            "Analyzes complex, multi-part enterprise questions.",
            "Decomposes prompt into 1-3 targeted sub-queries.",
            "Optimizes for multi-format retrieval across documents."
        ], RGBColor(56, 189, 248)),
        ("2. Retriever Agent", "Evidence Gathering", [
            "Tool: VectorSearchTool (ChromaDB top-k similarity).",
            "Tool: DocumentMetadataTool (extracts provenance).",
            "Deduplicates chunks across PDF, TXT, CSV, & Excel."
        ], RGBColor(99, 102, 241)),
        ("3. Reasoning Agent", "Context Synthesis", [
            "Reasons across heterogeneous multi-format chunks.",
            "Synthesizes grounded decision support answers.",
            "Enforces strict inline citations: [Source | Location]."
        ], RGBColor(16, 185, 129)),
        ("4. Validator Agent", "Hallucination Safeguard", [
            "Tool: HallucinationAuditTool (numerical cross-check).",
            "Audits response against source vector chunks.",
            "Assigns Verification Status badge: PASSED / FLAGGED."
        ], RGBColor(245, 158, 11))
    ]

    col_width = Inches(2.7)
    gap = Inches(0.3)
    left_margin = Inches(0.8)
    top_pos = Inches(2.0)
    card_height = Inches(4.8)

    for i, (title, subtitle, bullets, color) in enumerate(agents):
        x = left_margin + i * (col_width + gap)
        box = slide1.shapes.add_textbox(x, top_pos, col_width, card_height)
        btf = box.text_frame
        btf.word_wrap = True
        
        hp = btf.paragraphs[0]
        hp.text = title
        hp.font.size = Pt(18)
        hp.font.bold = True
        hp.font.color.rgb = color
        
        sp = btf.add_paragraph()
        sp.text = subtitle
        sp.font.size = Pt(13)
        sp.font.italic = True
        sp.font.color.rgb = RGBColor(100, 116, 139)
        sp.space_after = Pt(14)
        
        for b in bullets:
            bp = btf.add_paragraph()
            bp.text = f"• {b}"
            bp.font.size = Pt(12)
            bp.font.color.rgb = RGBColor(30, 41, 59)
            bp.space_after = Pt(8)

    # --- SLIDE 2: Ingestion, Guardrails, and Limitations ---
    slide2 = prs.slides.add_slide(blank_layout)

    tx_box2 = slide2.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.2))
    tf2 = tx_box2.text_frame
    tf2.word_wrap = True
    p_s2 = tf2.paragraphs[0]
    p_s2.text = "Ingestion Pipeline, Safety Guardrails & Engineering Limitations"
    p_s2.font.size = Pt(28)
    p_s2.font.bold = True
    p_s2.font.color.rgb = RGBColor(15, 23, 42)
    p_s2_sub = tf2.add_paragraph()
    p_s2_sub.text = "Complete Technical Specification and Mitigation Strategies"
    p_s2_sub.font.size = Pt(14)
    p_s2_sub.font.color.rgb = RGBColor(71, 85, 105)

    sections = [
        ("Multi-Format Ingestion (Tasks 3 & 4)", [
            "PDF: Page-level digital text extraction via pypdf.",
            "TXT: Unstructured system & incident log parsing (UTF-8).",
            "CSV: Row-to-text semantic record serializer via pandas.",
            "Excel (XLSX): Multi-sheet workbook parsing via openpyxl.",
            "Chunking: RecursiveCharacterTextSplitter (800 / 150 overlap)."
        ], Inches(0.8), Inches(2.0), Inches(5.6), Inches(2.3)),

        ("Reliability & Safety Guardrails (Task 9)", [
            "Input Sanitization: Blocks prompt injections via regex filter.",
            "Privacy Filter: Restricts queries for credentials/PII.",
            "Domain Safety: Refuses clinical medical prescription requests.",
            "File Security: Enforces 25MB limits & allowed extensions.",
            "Hallucination Auditing: Self-reflective LLM evaluation badge."
        ], Inches(6.9), Inches(2.0), Inches(5.6), Inches(2.3)),

        ("Key Limitations & Mitigations (Task 10 / Section 5)", [
            "1. Tabular vs. Semantic Search: Vector search cannot do dynamic SQL SUM/AVG across thousands of rows. Mitigation: Serialized headers into every row chunk to preserve schema context.",
            "2. Remote Model Lifecycle: Cloud endpoints retire older model checkpoints (410 Gone). Mitigation: Engineered automatic fallback cascades across active NVIDIA NIM and OpenAI models.",
            "3. Multi-Agent Latency: 4 sequential agent steps take 2-4s. Mitigation: Streamed live execution trace via Streamlit st.status for full user transparency."
        ], Inches(0.8), Inches(4.6), Inches(11.7), Inches(2.2))
    ]

    for stitle, sbullets, sx, sy, sw, sh in sections:
        sbox = slide2.shapes.add_textbox(sx, sy, sw, sh)
        stf = sbox.text_frame
        stf.word_wrap = True
        
        stp = stf.paragraphs[0]
        stp.text = stitle
        stp.font.size = Pt(17)
        stp.font.bold = True
        stp.font.color.rgb = RGBColor(30, 41, 59)
        stp.space_after = Pt(8)
        
        for sb in sbullets:
            sbp = stf.add_paragraph()
            sbp.text = sb if sb.startswith("1.") or sb.startswith("2.") or sb.startswith("3.") else f"• {sb}"
            sbp.font.size = Pt(11)
            sbp.font.color.rgb = RGBColor(71, 85, 105)
            sbp.space_after = Pt(4)

    prs.save(output_path)
    print(f"Generated PowerPoint presentation: {output_path}")


if __name__ == "__main__":
    build_presentation()
